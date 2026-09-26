"""
단기 근로자 인력공급 — 계획/실적표 생성 스크립트
실행:  python3 labor-kit/build_workbook.py
출력:  labor-kit/계획실적표_템플릿.xlsx

시트 구성
  안내        사용법과 색상 범례
  고객사      고객사 마스터 (차량지원 여부·단가)
  일일계획실적 날짜·고객사·직무·교대별 계획/실적 입력 (나머지는 자동 계산)
  월간요약    기준월 한 칸만 바꾸면 고객사별·일자별 실적이 자동 집계
  목록        드롭다운 목록 (직무·교대·지원형태)
"""
import datetime as dt
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).with_name("계획실적표_템플릿.xlsx")

FONT = "Arial"
NAVY = "1F3A5F"
INPUT_FILL = PatternFill("solid", fgColor="FFF4C2")   # 입력 칸
CALC_FILL = PatternFill("solid", fgColor="F2F4F7")    # 자동 계산 칸
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
TOTAL_FILL = PatternFill("solid", fgColor="DCE6F1")
THIN = Side(style="thin", color="C9CED6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

VND = '#,##0;(#,##0);"-"'
INT = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
DATE = "yyyy-mm-dd"

CLIENT_ROWS = 100      # 고객사 입력 가능 행 수
DAILY_ROWS = 2000      # 일일 입력 가능 행 수
SUMMARY_CLIENTS = 30   # 월간요약에 표시할 고객사 수


def f(bold=False, color="000000", size=10):
    return Font(name=FONT, bold=bold, color=color, size=size)


def header(ws, row, labels, widths=None):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font = f(True, "FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = 32
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def style_cell(c, fill, fmt=None, align="center"):
    c.font = f()
    c.fill = fill
    c.border = BORDER
    c.alignment = Alignment(horizontal=align, vertical="center")
    if fmt:
        c.number_format = fmt


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = f(True, NAVY, 16)
    ws["A2"] = sub
    ws["A2"].font = f(False, "595959", 9)


wb = Workbook()

# ------------------------------------------------------------------ 목록
lists = wb.active
lists.title = "목록"
list_data = {
    "A": ("직무 / Công việc", ["생산 라인", "포장", "검사(QC)", "물류·상하차", "창고 피킹", "청소·잡무", "봉제", "기타"]),
    "B": ("교대 / Ca", ["주간", "야간", "2교대", "3교대"]),
    "C": ("차량지원 / Hỗ trợ xe", ["Y", "N"]),
    "D": ("지원형태 / Hình thức", ["고객사 셔틀버스", "고객사 교통비 지급", "자사 차량 운행", "미지원(개별 출퇴근)"]),
}
for col, (head, items) in list_data.items():
    c = lists[f"{col}1"]
    c.value = head
    c.font = f(True, "FFFFFF")
    c.fill = HEAD_FILL
    for i, v in enumerate(items, start=2):
        lists[f"{col}{i}"] = v
        lists[f"{col}{i}"].font = f()
        lists[f"{col}{i}"].fill = INPUT_FILL
    lists.column_dimensions[col].width = 24
lists["F1"] = "노란 칸에 항목을 추가하면 드롭다운에 바로 반영됩니다 (각 열 30행까지)."
lists["F1"].font = f(False, "595959", 9)

# ------------------------------------------------------------------ 고객사
cl = wb.create_sheet("고객사", 0)
title(cl, "고객사 마스터 / Danh sách khách hàng",
      "노란 칸 입력. 코드는 중복 없이 입력하세요 (일일계획실적에서 코드로 연결됩니다).")
cl_headers = [
    "고객사코드\nMã KH", "고객사명\nTên khách hàng", "지역(성/시)\nTỉnh/TP", "업종\nNgành",
    "담당자\nNgười phụ trách", "연락처\nSĐT", "고객사 차량지원\n(Y/N)", "지원형태\nHình thức",
    "픽업지점\nĐiểm đón", "청구단가\n(VND/인·일)", "지급임금\n(VND/인·일)", "자사차량비\n(VND/대·일)",
    "계약시작\nBắt đầu", "계약종료\nKết thúc", "비고\nGhi chú",
]
header(cl, 4, cl_headers, [11, 22, 13, 12, 12, 14, 12, 18, 18, 14, 14, 14, 12, 12, 22])
clients = [
    ("C001", "예시) 한빛전자 비나", "박닌", "전자부품", "Nguyễn Văn A", "0912-000-001", "Y", "고객사 셔틀버스", "박닌 버스터미널", 420000, 330000, 0, dt.date(2026, 1, 1), dt.date(2026, 12, 31), "야간 할증 별도"),
    ("C002", "예시) 대성섬유 VN", "동나이", "봉제", "Trần Thị B", "0912-000-002", "N", "자사 차량 운행", "비엔호아 시장", 380000, 300000, 1200000, dt.date(2026, 3, 1), dt.date(2027, 2, 28), "29인승 버스 1대"),
    ("C003", "예시) 그린물류 하노이", "하노이", "물류", "Lê Văn C", "0912-000-003", "N", "미지원(개별 출퇴근)", "-", 400000, 315000, 0, dt.date(2026, 6, 1), dt.date(2026, 11, 30), ""),
    ("C004", "예시) 미래식품 빈즈엉", "빈즈엉", "식품", "Phạm Thị D", "0912-000-004", "Y", "고객사 교통비 지급", "-", 390000, 305000, 0, dt.date(2026, 5, 1), dt.date(2027, 4, 30), "교통비 1일 20,000 고객사 지급"),
]
first_c = 5
last_c = first_c + CLIENT_ROWS - 1
for r in range(first_c, last_c + 1):
    for col in range(1, len(cl_headers) + 1):
        fmt = VND if col in (10, 11, 12) else DATE if col in (13, 14) else None
        style_cell(cl.cell(row=r, column=col), INPUT_FILL, fmt, "left" if col in (2, 8, 9, 15) else "center")
for i, row in enumerate(clients):
    for col, v in enumerate(row, start=1):
        cl.cell(row=first_c + i, column=col, value=v)
cl.freeze_panes = "C5"
cl.cell(row=4, column=7).comment = Comment(
    "Y = 고객사가 출퇴근 차량(셔틀) 또는 교통비를 지원\nN = 미지원 → 자사 차량 운행 시 '자사차량비'가 비용으로 잡힙니다.", "템플릿")
cl.cell(row=4, column=12).comment = Comment(
    "자사 차량 1대를 하루 운행할 때 드는 비용(렌트+유류+기사). 일일계획실적의 '자사 차량운행(대)'와 곱해집니다.", "템플릿")

dv_yn = DataValidation(type="list", formula1="=목록!$C$2:$C$31", allow_blank=True)
dv_support = DataValidation(type="list", formula1="=목록!$D$2:$D$31", allow_blank=True)
cl.add_data_validation(dv_yn)
cl.add_data_validation(dv_support)
dv_yn.add(f"G{first_c}:G{last_c}")
dv_support.add(f"H{first_c}:H{last_c}")

CL_RANGE = f"고객사!$A${first_c}:$A${last_c}"


def lookup(col_letter, key):
    return f'INDEX(고객사!${col_letter}${first_c}:${col_letter}${last_c},MATCH({key},{CL_RANGE},0))'


# ------------------------------------------------------------------ 일일계획실적
d = wb.create_sheet("일일계획실적", 1)
title(d, "일일 계획 / 실적 — Kế hoạch & Thực tế hằng ngày",
      "노란 칸(날짜·고객사코드·직무·교대·계획·실적·자사차량·비고)만 입력하면 회색 칸은 자동 계산됩니다.")
d_headers = [
    "날짜\nNgày", "고객사코드\nMã KH", "고객사명\nTên KH", "직무\nCông việc", "교대\nCa",
    "계획인원\nKế hoạch", "실적인원\nThực tế", "달성률\nTỷ lệ đạt", "결원\nThiếu",
    "고객사 차량지원\nKH hỗ trợ xe", "자사 차량운행(대)\nXe công ty", "차량비용(VND)\nChi phí xe",
    "차량확인\nKiểm tra xe", "청구단가\nĐơn giá thu", "매출(VND)\nDoanh thu",
    "지급임금\nĐơn giá lương", "인건비(VND)\nChi phí lương", "이익(VND)\nLợi nhuận",
    "이익률\nTỷ suất", "비고\nGhi chú",
]
header(d, 4, d_headers, [12, 11, 22, 13, 8, 10, 10, 9, 8, 12, 12, 14, 11, 12, 15, 12, 15, 15, 9, 24])
first_d = 5
last_d = first_d + DAILY_ROWS - 1
input_cols = {1, 2, 4, 5, 6, 7, 11, 20}
for r in range(first_d, last_d + 1):
    b = f"$B{r}"
    formulas = {
        3: f'=IF({b}="","",IFERROR({lookup("B", b)},"코드 확인"))',
        8: f'=IF(OR({b}="",F{r}=0),"",G{r}/F{r})',
        9: f'=IF({b}="","",MAX(F{r}-G{r},0))',
        10: f'=IF({b}="","",IFERROR({lookup("G", b)},""))',
        12: f'=IF({b}="","",K{r}*IFERROR({lookup("L", b)},0))',
        13: f'=IF({b}="","",IF(AND(J{r}="N",G{r}>0,K{r}=0,IFERROR({lookup("H", b)},"")="자사 차량 운행"),"차량 미배정",""))',
        14: f'=IF({b}="","",IFERROR({lookup("J", b)},0))',
        15: f'=IF({b}="","",G{r}*N{r})',
        16: f'=IF({b}="","",IFERROR({lookup("K", b)},0))',
        17: f'=IF({b}="","",G{r}*P{r})',
        18: f'=IF({b}="","",O{r}-Q{r}-L{r})',
        19: f'=IF(OR({b}="",N(O{r})=0),"",R{r}/O{r})',
        # 보조열: 같은 날짜·고객사의 첫 행이면 1 (월간요약 운영일수 계산용)
        21: f'=IF(OR({b}="",$A{r}=""),0,IF(COUNTIFS($A${first_d}:$A{r},$A{r},$B${first_d}:$B{r},{b})=1,1,0))',
    }
    fmts = {1: DATE, 6: INT, 7: INT, 8: PCT, 9: INT, 11: INT, 12: VND, 14: VND, 15: VND, 16: VND, 17: VND, 18: VND, 19: PCT}
    for col in range(1, len(d_headers) + 1):
        c = d.cell(row=r, column=col)
        if col in formulas:
            c.value = formulas[col]
        style_cell(c, INPUT_FILL if col in input_cols else CALC_FILL, fmts.get(col),
                   "left" if col in (3, 20) else "center")
    d.cell(row=r, column=21, value=formulas[21]).font = f(False, "808080", 8)
d.freeze_panes = "D5"
d["U4"] = "보조(운영일)"
d["U4"].font = f(False, "808080", 8)
d.column_dimensions["U"].hidden = True
d.auto_filter.ref = f"A4:T{last_d}"

dv_code = DataValidation(type="list", formula1=f"={CL_RANGE}", allow_blank=True)
dv_job = DataValidation(type="list", formula1="=목록!$A$2:$A$31", allow_blank=True)
dv_shift = DataValidation(type="list", formula1="=목록!$B$2:$B$31", allow_blank=True)
dv_num = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                        error="0 이상의 정수를 입력하세요.", showErrorMessage=True)
dv_date = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                         error="날짜 형식(예: 2026-10-01)으로 입력하세요.", showErrorMessage=True)
for dv in (dv_code, dv_job, dv_shift, dv_num, dv_date):
    d.add_data_validation(dv)
dv_code.add(f"B{first_d}:B{last_d}")
dv_job.add(f"D{first_d}:D{last_d}")
dv_shift.add(f"E{first_d}:E{last_d}")
dv_num.add(f"F{first_d}:G{last_d}")
dv_num.add(f"K{first_d}:K{last_d}")
dv_date.add(f"A{first_d}:A{last_d}")

red_font = Font(name=FONT, color="C00000", bold=True)
d.conditional_formatting.add(
    f"H{first_d}:H{last_d}",
    FormulaRule(formula=[f'AND(ISNUMBER(H{first_d}),H{first_d}<0.9)'], font=red_font,
                fill=PatternFill("solid", fgColor="FDE2E2")))
d.conditional_formatting.add(
    f"M{first_d}:M{last_d}",
    CellIsRule(operator="equal", formula=['"차량 미배정"'], font=red_font,
               fill=PatternFill("solid", fgColor="FDE2E2")))
d.conditional_formatting.add(
    f"C{first_d}:C{last_d}",
    CellIsRule(operator="equal", formula=['"코드 확인"'], font=red_font))

# 예시 데이터: 2026년 9월 1~25일 (삭제 후 사용)
random.seed(7)
plan_rows = [("C001", "생산 라인", "주간", 40, 0), ("C001", "검사(QC)", "야간", 12, 0),
             ("C002", "봉제", "주간", 30, 1), ("C003", "물류·상하차", "주간", 18, 0),
             ("C004", "포장", "2교대", 22, 0)]
r = first_d
for day in range(1, 26):
    date = dt.date(2026, 9, day)
    if date.weekday() == 6:
        continue
    for code, job, shift, plan, vehicles in plan_rows:
        actual = max(plan - random.choice([0, 0, 0, 1, 1, 2, 3, 5]), 0)
        d.cell(row=r, column=1, value=date)
        d.cell(row=r, column=2, value=code)
        d.cell(row=r, column=4, value=job)
        d.cell(row=r, column=5, value=shift)
        d.cell(row=r, column=6, value=plan)
        d.cell(row=r, column=7, value=actual)
        d.cell(row=r, column=11, value=vehicles)
        d.cell(row=r, column=20, value="예시 데이터 — 삭제 후 사용" if r == first_d else None)
        r += 1

# ------------------------------------------------------------------ 월간요약
s = wb.create_sheet("월간요약", 2)
title(s, "월간 계획 / 실적 요약 — Báo cáo tháng",
      "C3 기준월만 바꾸면 전체가 다시 계산됩니다. 인원은 '연인원(인·일)' 기준입니다.")
s["A3"] = "기준월 / Tháng"
s["A3"].font = f(True)
s["C3"] = dt.date(2026, 9, 1)
style_cell(s["C3"], INPUT_FILL, "yyyy-mm")
s["C3"].font = f(True, "0000FF", 11)
s["C3"].comment = Comment("해당 월의 1일을 입력하세요 (예: 2026-10-01).", "템플릿")
s["E3"] = "월 시작"
s["F3"] = "=DATE(YEAR($C$3),MONTH($C$3),1)"
s["G3"] = "다음 달 시작"
s["H3"] = "=EDATE($F$3,1)"
for ref in ("F3", "H3"):
    style_cell(s[ref], CALC_FILL, DATE)
for ref in ("E3", "G3"):
    s[ref].font = f(False, "595959", 9)

s_headers = [
    "고객사코드\nMã KH", "고객사명\nTên KH", "차량지원\nHỗ trợ xe", "운영일수\nSố ngày",
    "계획 연인원\nKế hoạch", "실적 연인원\nThực tế", "달성률\nTỷ lệ đạt", "결원 합계\nThiếu",
    "매출(VND)\nDoanh thu", "인건비(VND)\nLương", "차량비(VND)\nChi phí xe", "이익(VND)\nLợi nhuận", "이익률\nTỷ suất",
]
header(s, 5, s_headers, [11, 24, 10, 10, 12, 12, 10, 10, 16, 16, 14, 16, 10])
DR = f"$A${first_d}:$A${last_d}"


def sumifs(col, key):
    rng = f"일일계획실적!${col}${first_d}:${col}${last_d}"
    return (f"SUMIFS({rng},일일계획실적!$B${first_d}:$B${last_d},{key},"
            f"일일계획실적!{DR},\">=\"&$F$3,일일계획실적!{DR},\"<\"&$H$3)")


first_s = 6
last_s = first_s + SUMMARY_CLIENTS - 1
for i in range(SUMMARY_CLIENTS):
    r = first_s + i
    src = first_c + i
    key = f"$A{r}"
    row = {
        1: f'=IF(고객사!$A${src}="","",고객사!$A${src})',
        2: f'=IF({key}="","",고객사!$B${src})',
        3: f'=IF({key}="","",고객사!$G${src})',
        4: f'=IF({key}="","",{sumifs("U", key)})',
        5: f'=IF({key}="","",{sumifs("F", key)})',
        6: f'=IF({key}="","",{sumifs("G", key)})',
        7: f'=IF(OR({key}="",N(E{r})=0),"",F{r}/E{r})',
        8: f'=IF({key}="","",{sumifs("I", key)})',
        9: f'=IF({key}="","",{sumifs("O", key)})',
        10: f'=IF({key}="","",{sumifs("Q", key)})',
        11: f'=IF({key}="","",{sumifs("L", key)})',
        12: f'=IF({key}="","",I{r}-J{r}-K{r})',
        13: f'=IF(OR({key}="",N(I{r})=0),"",L{r}/I{r})',
    }
    fmts = {4: INT, 5: INT, 6: INT, 7: PCT, 8: INT, 9: VND, 10: VND, 11: VND, 12: VND, 13: PCT}
    for col, v in row.items():
        c = s.cell(row=r, column=col, value=v)
        style_cell(c, CALC_FILL, fmts.get(col), "left" if col == 2 else "center")

tr = last_s + 1
s.cell(row=tr, column=1, value="합계 / Tổng")
for col in range(1, 14):
    c = s.cell(row=tr, column=col)
    style_cell(c, TOTAL_FILL, None, "center")
    c.font = f(True)
for col in (5, 6, 8, 9, 10, 11, 12):
    L = get_column_letter(col)
    c = s.cell(row=tr, column=col, value=f"=SUM({L}{first_s}:{L}{last_s})")
    c.number_format = VND if col >= 9 else INT
c = s.cell(row=tr, column=7, value=f'=IF(E{tr}=0,"",F{tr}/E{tr})')
c.number_format = PCT
c = s.cell(row=tr, column=13, value=f'=IF(I{tr}=0,"",L{tr}/I{tr})')
c.number_format = PCT
s.conditional_formatting.add(
    f"G{first_s}:G{tr}",
    FormulaRule(formula=[f'AND(ISNUMBER(G{first_s}),G{first_s}<0.95)'], font=red_font))

# 일자별 추이 (차트 원본)
dr0 = tr + 3
s.cell(row=dr0 - 1, column=1, value="일자별 추이 / Theo ngày").font = f(True, NAVY, 12)
header_row = dr0
for col, label in enumerate(["일자\nNgày", "날짜", "계획\nKế hoạch", "실적\nThực tế", "달성률\nTỷ lệ"], start=1):
    c = s.cell(row=header_row, column=col, value=label)
    c.font = f(True, "FFFFFF")
    c.fill = HEAD_FILL
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BORDER
s.row_dimensions[header_row].height = 32
for day in range(1, 32):
    r = header_row + day
    s.cell(row=r, column=1, value=day)
    s.cell(row=r, column=2, value=f'=IF({day}>DAY(EOMONTH($F$3,0)),"",$F$3+{day - 1})')
    s.cell(row=r, column=3, value=(f'=IF(B{r}="",0,SUMIFS(일일계획실적!$F${first_d}:$F${last_d},'
                                   f'일일계획실적!{DR},B{r}))'))
    s.cell(row=r, column=4, value=(f'=IF(B{r}="",0,SUMIFS(일일계획실적!$G${first_d}:$G${last_d},'
                                   f'일일계획실적!{DR},B{r}))'))
    s.cell(row=r, column=5, value=f'=IF(C{r}=0,"",D{r}/C{r})')
    for col, fmt in zip(range(1, 6), (None, DATE, INT, INT, PCT)):
        style_cell(s.cell(row=r, column=col), CALC_FILL, fmt)
last_day_row = header_row + 31

bar = BarChart()
bar.type = "col"
bar.title = "일자별 계획 vs 실적 (명)"
bar.y_axis.title = "인원"
bar.add_data(Reference(s, min_col=3, max_col=4, min_row=header_row, max_row=last_day_row), titles_from_data=True)
bar.set_categories(Reference(s, min_col=1, min_row=header_row + 1, max_row=last_day_row))
bar.series[0].graphicalProperties.solidFill = "B7C4D6"
bar.series[0].graphicalProperties.line.noFill = True
bar.series[1].graphicalProperties.solidFill = NAVY
bar.series[1].graphicalProperties.line.noFill = True
bar.gapWidth = 60
bar.height = 8.5
bar.width = 22
s.add_chart(bar, f"G{dr0 - 1}")
s.freeze_panes = "C6"

# ------------------------------------------------------------------ 안내
g = wb.create_sheet("안내", 0)
g.column_dimensions["A"].width = 3
g.column_dimensions["B"].width = 26
g.column_dimensions["C"].width = 90
g["B2"] = "인력공급 계획/실적표 사용법"
g["B2"].font = f(True, NAVY, 16)
g["B3"] = "Hướng dẫn sử dụng — Bảng kế hoạch & thực tế cung ứng lao động"
g["B3"].font = f(False, "595959", 10)
rows = [
    ("① 고객사", "고객사 코드·이름·단가·차량지원 여부(Y/N)를 한 번만 등록합니다. 지원형태가 '자사 차량 운행'이면 자사차량비(대·일)를 입력하세요."),
    ("② 일일계획실적", "매일 날짜 / 고객사코드(드롭다운) / 직무 / 교대 / 계획인원 / 실적인원을 입력합니다. 자사 차량을 운행했다면 대수를 입력하세요."),
    ("", "고객사명·달성률·결원·차량지원·매출·인건비·이익은 자동 계산됩니다. 달성률 90% 미만, '차량 미배정'은 빨간색으로 표시됩니다."),
    ("③ 월간요약", "C3 기준월(예: 2026-10-01)만 바꾸면 고객사별 연인원·달성률·매출·이익과 일자별 추이 차트가 갱신됩니다."),
    ("④ 목록", "직무·교대·지원형태 드롭다운 항목을 추가/수정할 수 있습니다."),
    ("예시 데이터", "고객사 C001~C004와 2026년 9월 일일 데이터는 예시입니다. 실제 사용 전 노란 칸의 예시 값을 지우세요 (수식 칸은 지우지 마세요)."),
    ("제안서 연결", "월간요약의 달성률·연인원을 제안서(labor-kit/proposal/data/company.json의 performance)에 옮기면 제안서 실적 차트에 반영됩니다."),
]
for i, (k, v) in enumerate(rows, start=5):
    g.cell(row=i, column=2, value=k).font = f(True)
    c = g.cell(row=i, column=3, value=v)
    c.font = f()
    c.alignment = Alignment(wrap_text=True, vertical="top")
    g.row_dimensions[i].height = 30
lg = 5 + len(rows) + 1
g.cell(row=lg, column=2, value="색상 범례").font = f(True, NAVY, 12)
for i, (fill, text) in enumerate([(INPUT_FILL, "입력 칸 — 직접 입력/수정"),
                                  (CALC_FILL, "자동 계산 — 수식이 들어 있으니 지우지 마세요"),
                                  (TOTAL_FILL, "합계")], start=lg + 1):
    g.cell(row=i, column=2).fill = fill
    g.cell(row=i, column=2).border = BORDER
    g.cell(row=i, column=3, value=text).font = f()
g.sheet_view.showGridLines = False

for ws in wb.worksheets:
    ws.sheet_properties.tabColor = {"안내": "7F7F7F", "고객사": "2E75B6", "일일계획실적": "C9A227",
                                    "월간요약": NAVY, "목록": "BFBFBF"}[ws.title]
wb.active = wb.sheetnames.index("일일계획실적")
wb.save(OUT)
print(f"saved {OUT}")
